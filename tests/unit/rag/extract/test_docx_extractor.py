"""
Unit tests for DocxExtractor.

Each test builds a small real .docx file in tmp_path with python-docx, so the
extractor runs against genuine document XML.

Error tests assert only the HTTP status code of the raised ApiException.
"""

import zipfile
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml import parse_xml

from src.exceptions.base import ApiException
from src.rag.services.extract.extractors import docx_extractor
from src.rag.services.extract.extractors.docx_extractor import DocxExtractor

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ---------------------------------------------------------------------------
# Fixtures and helpers
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def segment_target():
    """A large target, so segments only close at headings unless a test says so."""
    fake = SimpleNamespace(rag=SimpleNamespace(SEGMENT_TARGET_CHARS=10_000))
    with patch("src.rag.services.extract.extractors.extractor.settings", fake):
        yield fake


def save(document, tmp_path, name="doc.docx"):
    path = tmp_path / name
    document.save(path)
    return path


def extract_texts(path) -> list[str]:
    return [s.text for s in DocxExtractor().extract(path)]


def content_control(text: str):
    return parse_xml(
        f'<w:sdt xmlns:w="{W_NS}"><w:sdtContent><w:p><w:r><w:t>{text}</w:t></w:r>'
        f"</w:p></w:sdtContent></w:sdt>"
    )


# ---------------------------------------------------------------------------
# Paragraphs, headings and segments
# ---------------------------------------------------------------------------


class TestParagraphsAndHeadings:

    def test_paragraphs_are_joined_into_one_segment(self, tmp_path):
        doc = Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")

        segments = list(DocxExtractor().extract(save(doc, tmp_path)))

        assert len(segments) == 1
        assert segments[0].text == "First paragraph.\nSecond paragraph."
        assert segments[0].page is None
        assert segments[0].separator == "\n\n"

    def test_blank_paragraphs_are_skipped_and_text_is_stripped(self, tmp_path):
        doc = Document()
        doc.add_paragraph("   padded   ")
        doc.add_paragraph("")
        doc.add_paragraph(" \t ")
        doc.add_paragraph("next")

        assert extract_texts(save(doc, tmp_path)) == ["padded\nnext"]

    def test_headings_get_markdown_prefixes_and_start_new_segments(self, tmp_path):
        doc = Document()
        doc.add_heading("Main Title", 0)
        doc.add_paragraph("Intro.")
        doc.add_heading("Section", 1)
        doc.add_paragraph("Body.")
        doc.add_heading("Sub", 3)
        doc.add_paragraph("Detail.")

        assert extract_texts(save(doc, tmp_path)) == [
            "# Main Title\nIntro.",
            "# Section\nBody.",
            "### Sub\nDetail.",
        ]

    def test_heading_level_is_capped_at_six(self, tmp_path):
        doc = Document()
        doc.add_heading("Deep", 9)

        assert extract_texts(save(doc, tmp_path)) == ["###### Deep"]

    def test_text_before_the_first_heading_forms_its_own_segment(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Preamble.")
        doc.add_heading("Section", 1)
        doc.add_paragraph("Body.")

        assert extract_texts(save(doc, tmp_path)) == [
            "Preamble.",
            "# Section\nBody.",
        ]

    def test_segment_closes_once_it_reaches_the_target_size(
        self, tmp_path, segment_target
    ):
        segment_target.rag.SEGMENT_TARGET_CHARS = 20
        doc = Document()
        for text in ("a" * 15, "b" * 15, "c" * 15):
            doc.add_paragraph(text)

        assert extract_texts(save(doc, tmp_path)) == [
            "a" * 15 + "\n" + "b" * 15,
            "c" * 15,
        ]

    def test_empty_document_yields_no_segments(self, tmp_path):
        assert extract_texts(save(Document(), tmp_path)) == []

    def test_non_ascii_text_is_preserved(self, tmp_path):
        doc = Document()
        doc.add_paragraph("Café résumé 日本語 🙂")

        assert extract_texts(save(doc, tmp_path)) == ["Café résumé 日本語 🙂"]


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


class TestTables:

    def test_each_row_becomes_one_line_with_cells_joined_by_a_pipe(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for r, row in enumerate([("Name", "Age"), ("Ada", "36")]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value

        assert extract_texts(save(doc, tmp_path)) == ["Name | Age\nAda | 36"]

    def test_empty_cells_and_empty_rows_are_skipped(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "a"
        table.cell(0, 2).text = "c"
        table.cell(2, 1).text = "z"

        assert extract_texts(save(doc, tmp_path)) == ["a | c\nz"]

    def test_horizontally_merged_cell_is_read_once(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "wide"
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 2).text = "end"

        assert extract_texts(save(doc, tmp_path)) == ["wide | end"]

    def test_vertically_merged_cell_is_read_once_not_per_row(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        for r in range(3):
            table.cell(r, 1).text = f"r{r}"
        table.cell(0, 0).text = "tall"
        table.cell(0, 0).merge(table.cell(2, 0))

        text = extract_texts(save(doc, tmp_path))[0]

        assert text == "tall | r0\nr1\nr2"
        assert text.count("tall") == 1

    def test_nested_table_text_is_included(self, tmp_path):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "outer"
        inner = table.cell(0, 1).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "inner1"
        inner.cell(0, 1).text = "inner2"

        assert extract_texts(save(doc, tmp_path)) == ["outer | inner1 | inner2"]

    def test_table_stays_between_the_paragraphs_around_it(self, tmp_path):
        doc = Document()
        doc.add_paragraph("before")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "x"
        table.cell(0, 1).text = "y"
        doc.add_paragraph("after")

        assert extract_texts(save(doc, tmp_path)) == ["before\nx | y\nafter"]


# ---------------------------------------------------------------------------
# Content controls
# ---------------------------------------------------------------------------


class TestContentControls:

    def test_text_inside_a_content_control_is_included_in_order(self, tmp_path):
        doc = Document()
        doc.add_paragraph("before")
        doc.element.body.insert(1, content_control("inside control"))
        doc.add_paragraph("after")

        assert extract_texts(save(doc, tmp_path)) == ["before\ninside control\nafter"]


# ---------------------------------------------------------------------------
# Error handling
# ---------------------------------------------------------------------------


def corrupt_copy(source, tmp_path, *, drop=None, replace=None):
    """Copy a docx zip, dropping one part or replacing the bytes of one."""
    target = tmp_path / "corrupt.docx"
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w") as zout:
        for name in zin.namelist():
            if name == drop:
                continue
            data = replace[1] if replace and name == replace[0] else zin.read(name)
            zout.writestr(name, data)
    return target


class TestErrorHandling:

    def test_file_that_is_not_a_zip_raises_422(self, tmp_path):
        path = tmp_path / "fake.docx"
        path.write_bytes(b"plain text, not a docx")

        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(path))

        assert exc_info.value.status_code == 422

    def test_missing_file_raises_422(self, tmp_path):
        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(tmp_path / "missing.docx"))

        assert exc_info.value.status_code == 422

    def test_truncated_file_raises_422(self, tmp_path):
        doc = Document()
        doc.add_paragraph("text")
        raw = save(doc, tmp_path).read_bytes()
        path = tmp_path / "truncated.docx"
        path.write_bytes(raw[: len(raw) // 2])

        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(path))

        assert exc_info.value.status_code == 422

    def test_invalid_xml_in_the_document_part_raises_422(self, tmp_path):
        doc = Document()
        doc.add_paragraph("text")
        source = save(doc, tmp_path)
        broken = corrupt_copy(
            source, tmp_path, replace=("word/document.xml", b"<w:document><broken")
        )

        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(broken))

        assert exc_info.value.status_code == 422

    def test_missing_document_part_raises_422(self, tmp_path):
        doc = Document()
        doc.add_paragraph("text")
        source = save(doc, tmp_path)
        broken = corrupt_copy(source, tmp_path, drop="word/document.xml")

        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(broken))

        assert exc_info.value.status_code == 422

    def test_file_over_the_uncompressed_size_limit_raises_422(
        self, tmp_path, monkeypatch
    ):
        doc = Document()
        doc.add_paragraph("text")
        path = save(doc, tmp_path)
        monkeypatch.setattr(docx_extractor, "_MAX_UNCOMPRESSED_BYTES", 1000)

        with pytest.raises(ApiException) as exc_info:
            list(DocxExtractor().extract(path))

        assert exc_info.value.status_code == 422
